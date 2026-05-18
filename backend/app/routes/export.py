"""REST API routes for exporting system data and import templates"""

import csv
import io

from app import db
from app.models import ExamAttempt, ExamSession, School, Student, Subject, User
from flask import jsonify, request, send_file
from flask_jwt_extended import get_jwt_identity, jwt_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import export_bp


def get_school_id_for_user(user):
    """Lấy ID trường học của user hiện tại"""
    return getattr(user, "school_id", None)


# ========================================================
# 1. API TẢI CÁC FILE CSV MẪU (DÙNG ĐỂ IMPORT)
# ========================================================


@export_bp.route("/template/<string:template_type>", methods=["GET"])
def download_template(template_type):
    """
    Cung cấp file CSV mẫu chuẩn để Giáo viên/QTV điền dữ liệu import
    template_type: 'student', 'question_mc', 'question_tf', 'question_sa'
    """
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)

    if template_type == "student":
        filename = "Mau_Nhap_Hoc_Sinh.csv"
        writer.writerow(
            [
                "cccd",
                "full_name",
                "gender",
                "date_of_birth",
                "address",
                "phone",
                "class_name",
                "tuchon1",
                "tuchon2",
            ]
        )
        writer.writerow(
            [
                "'012345678912",
                "Nguyễn Văn A",
                "Nam",
                "'2008-01-15",
                "123 Đường Trần Phú",
                "'0901234567",
                "12A1",
                "VAT_LI",
                "MT",
            ]
        )
        writer.writerow(
            [
                "'012345678913",
                "Trần Thị B",
                "Nữ",
                "'2008-11-20",
                "456 Đường Lê Lợi",
                "'0987654321",
                "12A2",
                "HOA_HO",
                "SINH_H",
            ]
        )

    elif template_type == "question_mc":
        filename = "Mau_Cau_Hoi_Trac_Nghiem.csv"
        writer.writerow(
            ["NoiDung", "Diem", "A", "B", "C", "D", "DapAnDung", "DinhHuong"]
        )
        writer.writerow(
            [
                "Đáp án nào dưới đây chỉ thủ đô của Việt Nam?",
                "0,25",
                "Hà Nội",
                "TP.HCM",
                "Đà Nẵng",
                "Huế",
                "A",
                "Chung",
            ]
        )

    elif template_type == "question_tf":
        filename = "Mau_Cau_Hoi_Dung_Sai.csv"
        writer.writerow(
            [
                "NoiDung",
                "Diem",
                "Y_a",
                "DS_a",
                "Y_b",
                "DS_b",
                "Y_c",
                "DS_c",
                "Y_d",
                "DS_d",
                "DinhHuong",
            ]
        )
        writer.writerow(
            [
                "Các nhận định sau đây đúng hay sai?",
                "1,0",
                "Hà Nội là thủ đô của Việt Nam",
                "Đúng",
                "Mặt trời mọc hướng Tây",
                "Sai",
                "$log(10) = 1$",
                "Đúng",
                "Nước sôi ở 50 độ C",
                "Sai",
                "Chung",
            ]
        )

    elif template_type == "question_sa":
        filename = "Mau_Cau_Hoi_Tra_Loi_Ngan.csv"
        writer.writerow(["NoiDung", "Diem", "DapAn", "DinhHuong"])
        writer.writerow(["Giải phương trình x - 2 = 0", "0,5", "2", "Chung"])

    else:
        return jsonify({"error": "Loại file mẫu không hợp lệ"}), 400

    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        mimetype="text/csv",
        as_attachment=True,
        download_name=filename,
    )


# ========================================================
# 2. API XUẤT DỮ LIỆU HỆ THỐNG (HỌC SINH / KẾT QUẢ)
# ========================================================
@export_bp.route(
    "/data/<string:entity_type>/<string:export_type>", methods=["POST", "GET"]
)
@jwt_required()
def export_system_data(entity_type, export_type):
    """
    Xuất dữ liệu học sinh hoặc kết quả ra file Excel/CSV
    QTV: Thấy toàn hệ thống (có ID Tỉnh, Phường/Xã)
    GVQL: Chỉ thấy dữ liệu của trường mình
    """
    try:
        current_user = User.query.get(get_jwt_identity())
        if not current_user:
            return jsonify({"error": "Người dùng không tồn tại"}), 401

        req_data = request.get_json() if request.is_json else {}
        headers = []
        data_rows = []
        filename = f"export_{entity_type}"

        # -----------------------------------------------
        # A. XUẤT DANH SÁCH HỌC SINH
        # -----------------------------------------------
        if entity_type == "students":
            if current_user.role == "admin":
                headers = [
                    "ID Tỉnh/Thành",
                    "ID Phường/Xã",
                    "Tỉnh/Thành",
                    "Phường/Xã",
                    "Trường Học",
                    "Mã Học Sinh",
                    "Họ và Tên",
                    "Mã Định Danh",
                    "Giới Tính",
                    "Ngày Sinh",
                    "Lớp",
                    "Địa Chỉ Chi Tiết",
                ]
                filename = "Danh_Sach_Hoc_Sinh_Toan_He_Thong"

                try:
                    students = (
                        db.session.query(
                            Student,
                            School) .outerjoin(
                            School,
                            Student.school_id == School.school_id) .order_by(
                            Student.province_id,
                            Student.ward_id,
                            School.school_name,
                            Student.class_name,
                            Student.full_name,
                        ) .all())
                except Exception:
                    db.session.rollback()
                    students = (
                        db.session.query(
                            Student,
                            School) .outerjoin(
                            School,
                            Student.school_id == School.school_id) .order_by(
                            School.school_name,
                            Student.class_name,
                            Student.full_name) .all())

                for stu, sch in students:
                    prov_id = getattr(
                        stu, "province_id", getattr(sch, "province_id", "")
                    )
                    ward_id = getattr(
                        stu, "ward_id", getattr(
                            sch, "ward_id", ""))
                    prov_name = (
                        getattr(stu.province, "name", prov_id)
                        if hasattr(stu, "province")
                        else prov_id
                    )
                    dist_name = (
                        getattr(
                            stu.district,
                            "name",
                            getattr(
                                stu,
                                "district_id",
                                "")) if hasattr(
                            stu,
                            "district") else "")

                    data_rows.append([prov_id,
                                      ward_id,
                                      prov_name,
                                      dist_name,
                                      sch.school_name if sch else "Chưa gán trường",
                                      stu.student_code or "",
                                      stu.full_name or "",
                                      stu.cccd or "",
                                      "Nam" if stu.gender == "male" else "Nữ",
                                      str(stu.date_of_birth) if stu.date_of_birth else "",
                                      stu.class_name or "",
                                      stu.address or "",
                                      ])

            elif current_user.role == "teacher":
                school_id = get_school_id_for_user(current_user)
                headers = [
                    "Mã Học Sinh",
                    "Họ và Tên",
                    "CCCD/Mã Định Danh",
                    "Giới Tính",
                    "Ngày Sinh",
                    "Lớp Học",
                    "Số Điện Thoại",
                    "Địa Chỉ",
                ]
                filename = "Danh_Sach_Hoc_Sinh"

                query = Student.query.filter_by(school_id=school_id)
                if req_data.get("class_name"):
                    query = query.filter_by(class_name=req_data["class_name"])
                    filename += f"_{req_data['class_name']}"

                students = query.order_by(
                    Student.class_name, Student.full_name).all()
                for s in students:
                    data_rows.append(
                        [
                            s.student_code or "",
                            s.full_name or "",
                            s.cccd or "",
                            "Nam" if s.gender == "male" else "Nữ",
                            str(s.date_of_birth) if s.date_of_birth else "",
                            s.class_name or "",
                            s.phone or "",
                            s.address or "",
                        ]
                    )

        # -----------------------------------------------
        # B. XUẤT KẾT QUẢ KỲ THI
        # -----------------------------------------------
        elif entity_type == "results":
            session_id = req_data.get("exam_session_id") or request.args.get(
                "exam_session_id"
            )
            subject_id = req_data.get(
                "subject_id") or request.args.get("subject_id")

            if not session_id:
                return jsonify(
                    {"error": "Vui lòng chọn kỳ thi để xuất kết quả"}), 400

            headers = [
                "Số Thứ Tự",
                "Mã Học Sinh",
                "Họ và Tên",
                "Mã định danh",
                "Ngày sinh",
                "Giới tính",
                "Lớp",
                "Trường",
                "Môn Thi",
                "Điểm Số",
                "Trạng Thái",
            ]
            filename = f"Ket_Qua_Thi_Ky_Thi_{session_id}"

            query = (
                db.session.query(
                    ExamAttempt,
                    Student,
                    Subject) .join(
                    Student,
                    ExamAttempt.student_id == Student.student_id) .outerjoin(
                    Subject,
                    ExamAttempt.subject_id == Subject.subject_id) .filter(
                    ExamAttempt.exam_session_id == int(session_id)))

            if current_user.role == "teacher":
                query = query.filter(
                    Student.school_id == get_school_id_for_user(current_user)
                )

            if subject_id:
                query = query.filter(ExamAttempt.subject_id == int(subject_id))

            attempts = query.order_by(
                Student.class_name, Student.full_name).all()

            for idx, (att, stu, subj) in enumerate(attempts, start=1):
                data_rows.append([idx,
                                  stu.student_code or "",
                                  stu.full_name,
                                  stu.cccd,
                                  str(stu.date_of_birth) if stu.date_of_birth else "",
                                  "Nam" if stu.gender == "male" else "Nữ",
                                  stu.class_name or "",
                                  stu.school.name if stu.school else "",
                                  subj.subject_name if subj else f"Mã: {att.subject_id}",
                                  float(att.total_score) if att.total_score is not None else 0.0,
                                  ("Đã hoàn thành" if att.status in ["graded",
                                                                     "completed"] 
                                                    else "Chưa nộp bài hoặc bài nộp không hợp lệ"),
                                  ])

        # -----------------------------------------------
        # XUẤT KẾT QUẢ THI CỦA TOÀN TRƯỜNG (DÀNH CHO GVQL)
        # -----------------------------------------------
        elif entity_type == "school_results":

            headers = [
                "Số Thứ Tự",
                "Mã Định Danh",
                "Họ và Tên",
                "Ngày Sinh",
                "Giới Tính",
                "Lớp",
                "Mã Môn Thi",
                "Tên Môn Thi",
                "Điểm",
                "Kỳ Thi",
                "Trạng Thái",
            ]
            filename = "Ket_Qua_Thi_Cua_Truong"

            query = (
                db.session.query(
                    ExamAttempt,
                    Student,
                    Subject,
                    ExamSession) .join(
                    Student,
                    ExamAttempt.student_id == Student.student_id) .outerjoin(
                    Subject,
                    ExamAttempt.subject_id == Subject.subject_id) .outerjoin(
                    ExamSession,
                    ExamAttempt.exam_session_id == ExamSession.exam_session_id,
                ))

            if current_user.role == "teacher":
                school_id = get_school_id_for_user(current_user)
                query = query.filter(Student.school_id == school_id)
            else:
                return (
                    jsonify(
                        {
                            "error": "Chỉ giáo viên quản lý mới được phép xuất dữ liệu này"
                        }
                    ),
                    403,
                )

            session_id = req_data.get("exam_session_id") or request.args.get(
                "exam_session_id"
            )
            if session_id:
                query = query.filter(
                    ExamAttempt.exam_session_id == int(session_id))
                filename += f"_KyThi_{session_id}"

            results = query.order_by(
                Student.class_name, Student.full_name).all()

            for idx, (att, stu, subj, sess) in enumerate(results, start=1):
                data_rows.append([idx,
                                  stu.cccd or "",
                                  stu.full_name or "",
                                  str(stu.date_of_birth) if stu.date_of_birth else "",
                                  "Nam" if stu.gender == "male" else "Nữ",
                                  stu.class_name or "",
                                  subj.subject_code if subj else "",
                                  subj.subject_name if subj else "",
                                  float(att.total_score) if att.total_score is not None else 0.0,
                                  sess.session_name if sess else "",
                                  ("Đã hoàn thành" if att.status in ["graded",
                                                                     "completed"] 
                                                    else "Chưa nộp bài hoặc bài nộp không hợp lệ"),
                                  ])

        else:
            return jsonify({"error": "Loại dữ liệu không hợp lệ"}), 400

        # ========================================================
        # ĐÓNG GÓI VÀ SINH FILE (XLSX / CSV)
        # ========================================================
        if export_type == "csv":
            output = io.StringIO()
            output.write("\ufeff")
            writer = csv.writer(
                output,
                delimiter=";",
                quoting=csv.QUOTE_MINIMAL)
            writer.writerow(headers)
            writer.writerows(data_rows)
            output.seek(0)
            return send_file(
                io.BytesIO(output.getvalue().encode("utf-8-sig")),
                mimetype="text/csv",
                as_attachment=True,
                download_name=f"{filename}.csv",
            )

        if export_type == "xlsx":
            wb = Workbook()
            ws = wb.active
            ws.title = "Dữ Liệu Hệ Thống"
            ws.views.sheetView[0].showGridLines = True

            font_header = Font(
                name="Arial",
                size=11,
                bold=True,
                color="FFFFFF")
            fill_header = PatternFill(
                start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
            )  # Xanh biển
            font_body = Font(name="Arial", size=11)
            thin_border = Border(
                left=Side(style="thin", color="E2E8F0"),
                right=Side(style="thin", color="E2E8F0"),
                top=Side(style="thin", color="E2E8F0"),
                bottom=Side(style="thin", color="E2E8F0"),
            )

            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = Alignment(
                    horizontal="center", vertical="center")
                cell.border = thin_border

            for r_data in data_rows:
                ws.append(r_data)
                curr_row = ws.max_row
                for col_idx in range(1, len(r_data) + 1):
                    cell = ws.cell(row=curr_row, column=col_idx)
                    cell.font = font_body
                    cell.border = thin_border
                    if isinstance(r_data[col_idx - 1], (int, float)):
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(
                            horizontal="left", wrap_text=True)

            for col in ws.columns:
                max_len = max((len(str(cell.value))
                               for cell in col if cell.value), default=0)
                ws.column_dimensions[col[0].column_letter].width = max(
                    max_len + 2, 12)

            file_stream = io.BytesIO()
            wb.save(file_stream)
            file_stream.seek(0)

            return send_file(
                file_stream,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name=f"{filename}.xlsx",
            )

        return jsonify({"error": "Định dạng file không được hỗ trợ"}), 400

    except Exception as e:
        db.session.rollback()
        return jsonify(
            {"error": f"Lỗi hệ thống khi trích xuất dữ liệu: {str(e)}"}), 500
